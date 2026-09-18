# coding='utf-8'
import sys

import time

from pathlib import Path

sys.path.append(str(Path(__file__).parent.parent))


from newPlayWright import PlayWright, logger
from openpyxl import load_workbook


def login(username, passwd):
    try:
        logger.info(f'【{username}】账号开始登录')
        url = 'https://www.moneyga.me/home/index/requirelogin'
        PlayWright.goto(url)

        username_ele = '//div[@style="display: block;"]//input[@placeholder="Username"]'
        PlayWright.input(username_ele, username)
        PlayWright.input('//div[@style="display: block;"]//input[@placeholder="Password"]', passwd)
        PlayWright.click('//button[text()="Enter Game"]')
        return True
    except Exception as e:
        logger.error(f'【{username}】账号登录失败：{e}')
        return False


def buy(username):
    try:
        logger.info(f'【{username}】账号开始购买体力')
        shops = [22, 152]
        for shop in shops:
            PlayWright.goto(f'https://www.moneyga.me/map/tile/{shop}')
            PlayWright.input('//input[@name="amount_3"]', '2')
            PlayWright.click('//button[text()="Buy Products"]')
            break
        PlayWright.goto('https://www.moneyga.me/dashboard/assetsdebts')
        PlayWright.click('(//div[text()="Use"])[1]')
        PlayWright.click('//div[text()="Confirm"]')
        success_ele = '//div[contains(text(), "You consumed a Food Package")]'
        time.sleep(1)
        status = PlayWright.get_count(success_ele)
        return status
    except Exception as e:
        logger.error(f'【{username}】账号购买体力失败：{e}')
        return False


def work(username):
    try:
        logger.info(f'【{username}】账号开始工作')
        PlayWright.goto('https://www.moneyga.me/dashboard/companies')
        PlayWright.click('(//td[text()="Worker at "]/a)[1]')
        PlayWright.click('//a[text()="Begin Work"]')

        for roll in range(10):
            PlayWright.click('//button[text()="Work Now"]')
            if not PlayWright.get_count('//a[contains(text(),"Click here to work your next job")]'):
                return True
            PlayWright.click('//a[contains(text(),"Click here to work your next job")]')

        logger.error(f'【{username}】账号工作失败')
        return False

    except Exception as e:
        logger.error(f'【{username}】账号工作失败:{e}')
        return False


def run_account(username, passwd):
    """单个账号执行：登录 -> 购买 -> 工作，返回写入备注的结果文本。"""
    logger.info(f'\n开始执行【{username}】账号....')
    money = login(username, passwd)
    if not money:
        return '登录失败'
    bought = buy(username)
    if not bought:
        return '购买失败'
    worked = work(username)
    return '成功' if worked else '失败'


def execute(file=None):
    """读取 xlsx，从第3行起按 账号/密码 执行 登录 -> 购买 -> 工作，
    并把结果写入该行“备注”列（E列）。"""

    account_col = 2  # B：账号
    pwd_col = 4  # D：密码
    remark_col = 5  # E：备注
    start_row = 3  # 数据起始行

    wb = load_workbook(file)
    ws = wb.active

    # 只收集账号非空的有效数据行，忽略尾部（及中间）空行
    data_rows = []
    for row in range(start_row, ws.max_row + 1):
        username = ws.cell(row=row, column=account_col).value
        passwd = ws.cell(row=row, column=pwd_col).value
        if username is None or passwd is None or str(username).strip() == '':
            continue
        data_rows.append((row, str(username).strip(), str(passwd).strip()))

    logger.info(f'开始执行,共{len(data_rows)}行数据')
    for row, username, passwd in data_rows:
        import time
        s1= time.time()
        try:
            remark = run_account(username, passwd)
        except Exception as e:
            remark = f'执行异常：{e}'
        ws.cell(row=row, column=remark_col).value = remark
        wb.save(file)  # 每行及时回写，避免中途中断丢失结果
        logger.info(f'结果：【{username}】{remark}')
        s2 = time.time()
        print(f'时间：{s2-s1}')
        PlayWright.clear_cookie()


if __name__ == '__main__':
    file = r'.\草药员工.xlsx'
    execute(file)

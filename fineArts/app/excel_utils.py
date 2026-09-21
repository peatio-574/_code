# -*- coding: utf-8 -*-
"""Excel 导入模板生成与解析"""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

# (中文表头, 模型字段, 是否必填, 列宽)
COLUMNS = [
    ('姓名', 'name', True, 12),
    ('发证日期', 'issue_date', True, 14),
    ('考级专业', 'major', True, 12),
    ('起止级', 'level_range', False, 10),
    ('等级', 'level', True, 8),
    ('报考考级机构', 'org', False, 20),
    ('证件号', 'id_card', True, 24),
    ('国籍', 'nationality', False, 10),
    ('民族', 'ethnicity', False, 10),
    ('出生日期', 'birth_date', False, 14),
    ('姓名全拼', 'name_pinyin', False, 18),
    ('证书编号', 'cert_no', True, 24),
]

HEADERS = [c[0] for c in COLUMNS]
FIELD_BY_HEADER = {c[0]: c[1] for c in COLUMNS}
REQUIRED_FIELDS = {c[1] for c in COLUMNS if c[2]}

EXAMPLE_ROW = {
    'name': '王沐凡',
    'issue_date': '2026-06-01',
    'major': '漫画',
    'level_range': '1-9',
    'level': '2',
    'org': '中国美术学院',
    'id_card': '610125201806120077',
    'nationality': '中国',
    'ethnicity': '汉族',
    'birth_date': '2018-06',
    'name_pinyin': 'Wang Mu Fan',
    'cert_no': '007203426102006880',
}

THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def build_template():
    """生成导入模板 xlsx（Sheet1 仅表头，便于直接填写；示例与说明放在第二个 Sheet）"""
    wb = Workbook()
    ws = wb.active
    ws.title = '证书导入模板'

    header_font = Font(color='000000', bold=True, size=11)

    for idx, (title, _field, required, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=title + ('*' if required else ''))
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = 'A2'

    # 第二个 Sheet：填写说明 + 示例
    ws2 = wb.create_sheet('填写说明')
    ws2.cell(row=1, column=1, value='填写说明：').font = Font(bold=True, color='333333')
    notes = [
        '1. 请在“证书导入模板”工作表中从第 2 行开始填写数据，带 * 的列为必填项。',
        '2. 证件号、证书编号、出生日期请在单元格中设置为“文本”格式，避免前导 0 丢失。',
        '3. 证书编号不可重复，重复数据上传时会自动跳过并提示。',
        '4. 下方为一行示例数据，仅供参考，请勿直接上传本说明工作表。',
    ]
    for i, text in enumerate(notes, start=2):
        ws2.cell(row=i, column=1, value=text).font = Font(color='606266', size=10)

    start = len(notes) + 3
    for idx, (title, field, _r, _w) in enumerate(COLUMNS, start=1):
        head = ws2.cell(row=start, column=idx, value=title)
        head.font = Font(bold=True)
        head.border = BORDER
        cell = ws2.cell(row=start + 1, column=idx, value=EXAMPLE_ROW.get(field, ''))
        cell.border = BORDER
        if field in ('id_card', 'cert_no', 'birth_date', 'issue_date'):
            cell.number_format = '@'
        ws2.column_dimensions[get_column_letter(idx)].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _find_header_row(ws, max_scan=5):
    for row in range(1, min(ws.max_row, max_scan) + 1):
        values = [_cell_text(ws.cell(row=row, column=c).value).replace('*', '').strip()
                  for c in range(1, len(HEADERS) + 1)]
        if '姓名' in values and '证书编号' in values:
            return row
    return None


def parse_workbook(file_stream):
    """解析导入的 xlsx，逐行校验必填字段。

    返回 (rows, fatal_error)
      rows: [{'_row': 行号, 'name': ..., ..., 'errors': [..]}]
      fatal_error: 字符串或 None（如文件格式错误）
    """
    try:
        if hasattr(file_stream, 'read'):
            stream = io.BytesIO(file_stream.read())
        else:
            stream = file_stream
        wb = load_workbook(stream, data_only=True)
    except Exception as exc:
        return [], '无法读取文件，请上传 .xlsx 格式的 Excel 文件（%s）' % exc

    ws = wb.active
    header_row = _find_header_row(ws)
    if not header_row:
        return [], '未找到表头，请使用模板文件（需包含“姓名”“证书编号”等列）'

    header_map = {}
    for col in range(1, ws.max_column + 1):
        title = _cell_text(ws.cell(row=header_row, column=col).value).replace('*', '')
        if title in FIELD_BY_HEADER:
            header_map[FIELD_BY_HEADER[title]] = col

    missing = REQUIRED_FIELDS - set(header_map.keys())
    if missing:
        names = [c[0] for c in COLUMNS if c[1] in missing]
        return [], '缺少必填列：%s' % '、'.join(names)

    rows = []
    for row in range(header_row + 1, ws.max_row + 1):
        data = {'_row': row, 'errors': []}
        empty = True
        for field, col in header_map.items():
            value = _cell_text(ws.cell(row=row, column=col).value)
            data[field] = value
            if value:
                empty = False
        if empty:
            continue
        for field in REQUIRED_FIELDS:
            if not data.get(field):
                label = [c[0] for c in COLUMNS if c[1] == field][0]
                data['errors'].append('%s不能为空' % label)
        rows.append(data)

    if not rows:
        return [], '未读取到有效数据行'

    return rows, None

# -*- coding: utf-8 -*-
"""后台管理：证书管理"""
import base64
import io
import os
import re
from datetime import datetime

from flask import (current_app, flash, redirect, render_template, request,
                   send_file, url_for)
from flask_login import current_user, login_required, login_user, logout_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ..api import fail, ok, pagination_dict
from ..excel_utils import COLUMNS, build_template, parse_workbook
from ..models import db, Admin, Certificate, Registration, Institution
from ..qr import generate_qr_png
from ..tokens import make_qr_token
from . import admin_bp

FIELDS = [c[1] for c in COLUMNS]


# --------------------------------------------------------------------------- #
# 登录 / 退出
# --------------------------------------------------------------------------- #
def _home_endpoint(user):
    """按角色返回登录后的首页：受限账号进入报名管理。"""
    if getattr(user, 'role', 'admin') == 'user':
        return url_for('admin.registrations')
    return url_for('admin.certificates')


@admin_bp.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(_home_endpoint(current_user))
    if request.method == 'POST':
        username = (request.form.get('username') or '').strip()
        password = request.form.get('password') or ''
        user = Admin.query.filter_by(username=username).first()
        if user and user.verify_password(password):
            login_user(user)
            return redirect(request.args.get('next') or _home_endpoint(user))
        flash('用户名或密码错误', 'error')
    return render_template('admin/login.html')


@admin_bp.before_request
def _restrict_limited_account():
    """受限账号（role=user）仅可访问报名管理、机构管理相关页面与接口。"""
    if not current_user.is_authenticated:
        return None
    if getattr(current_user, 'role', 'admin') != 'user':
        return None

    path = request.path
    allowed_prefixes = (
        '/admin/registrations',
        '/admin/registration/',
        '/admin/api/registrations',
        '/admin/api/registration/',
        '/admin/institutions',
        '/admin/api/institution',
        '/admin/institution/',
        '/admin/password',
        '/admin/logout',
    )
    if path.startswith(allowed_prefixes):
        return None
    if path in ('/admin', '/admin/'):
        return redirect(url_for('admin.registrations'))
    if path.startswith('/admin/api/'):
        return fail('无权访问', status=403)
    flash('无权访问该页面', 'error')
    return redirect(url_for('admin.registrations'))


@admin_bp.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('admin.login'))


@admin_bp.route('/password', methods=['GET', 'POST'])
@login_required
def change_password():
    """修改当前管理员密码"""
    if request.method == 'POST':
        old_password = request.form.get('old_password') or ''
        new_password = request.form.get('new_password') or ''
        confirm_password = request.form.get('confirm_password') or ''

        if not current_user.verify_password(old_password):
            flash('原密码错误', 'error')
        elif len(new_password) < 6:
            flash('新密码至少 6 位', 'error')
        elif new_password != confirm_password:
            flash('两次输入的新密码不一致', 'error')
        else:
            current_user.set_password(new_password)
            db.session.commit()
            flash('密码修改成功', 'success')
            return redirect(url_for('admin.change_password'))

    return render_template('admin/password.html')


# --------------------------------------------------------------------------- #
# 工具函数
# --------------------------------------------------------------------------- #
def _cert_no_exists(cert_no, exclude_id=None):
    if not cert_no:
        return False
    query = Certificate.query.filter(Certificate.cert_no == cert_no,
                                     Certificate.is_deleted.is_(False))
    if exclude_id:
        query = query.filter(Certificate.id != exclude_id)
    return query.first() is not None


# --------------------------------------------------------------------------- #
# 证书列表
# --------------------------------------------------------------------------- #
@admin_bp.route('/', methods=['GET'])
@admin_bp.route('/certificates', methods=['GET'])
@login_required
def certificates():
    """证书管理页面（固定 HTML，数据由 /admin/api/certificates 填充）"""
    majors = [row[0] for row in db.session.query(Certificate.major)
              .filter(Certificate.major.isnot(None), Certificate.major != '')
              .distinct().all()]
    return render_template('admin/certificates.html', majors=sorted(majors))


def _serialize(cert):
    data = cert.to_dict()
    data['is_deleted'] = cert.is_deleted
    return data


@admin_bp.route('/api/certificates')
def api_certificates():
    """证书列表接口（统一 JSON 格式：success + pagination + data）"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)

    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', current_app.config['PER_PAGE'], type=int)
    keyword = (request.args.get('q') or '').strip()
    major = (request.args.get('major') or '').strip()

    query = Certificate.query.filter(Certificate.is_deleted.is_(False))

    if keyword:
        like = '%%%s%%' % keyword
        query = query.filter(db.or_(Certificate.name.like(like),
                                    Certificate.cert_no.like(like),
                                    Certificate.id_card.like(like)))
    if major:
        query = query.filter(Certificate.major == major)

    pagination = query.order_by(Certificate.issue_date.desc(),
                                Certificate.updated_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False)
    data = [_serialize(c) for c in pagination.items]
    return ok(data, pagination=pagination_dict(pagination))


# --------------------------------------------------------------------------- #
# 校验
# --------------------------------------------------------------------------- #
def _validate(data):
    errors = []
    for title, field, required, _w in COLUMNS:
        if required and not data.get(field):
            errors.append('%s不能为空' % title)
    return errors


# --------------------------------------------------------------------------- #
# 软删除
# --------------------------------------------------------------------------- #
@admin_bp.route('/certificate/<int:cid>/delete', methods=['POST'])
@login_required
def certificate_delete(cid):
    cert = Certificate.query.get_or_404(cid)
    cert.is_deleted = True
    cert.deleted_at = datetime.now()
    db.session.commit()
    flash('已删除', 'success')
    return redirect(request.referrer or url_for('admin.certificates'))


@admin_bp.route('/api/certificate', methods=['POST'])
def api_certificate_create():
    """新增证书（弹窗用）"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    payload = request.get_json(silent=True) or request.form
    data = {field: (payload.get(field) or '').strip() for field in FIELDS}

    errors = _validate(data)
    if not errors and _cert_no_exists(data['cert_no']):
        errors.append('证书编号已存在：%s' % data['cert_no'])
    if errors:
        return fail('；'.join(errors), errors=errors)

    cert = Certificate(**data)
    db.session.add(cert)
    db.session.commit()
    return ok(_serialize(cert), message='新增成功')


@admin_bp.route('/api/certificate/<int:cid>', methods=['GET'])
def api_certificate_get(cid):
    """获取单条证书（编辑弹窗用）"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    cert = Certificate.query.get_or_404(cid)
    return ok(_serialize(cert))


@admin_bp.route('/api/certificate/<int:cid>/qrcode', methods=['GET'])
def api_certificate_qrcode(cid):
    """生成证书二维码（中间带 logo、底部带链接），返回详情链接与 data URI 图片"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    cert = Certificate.query.get_or_404(cid)

    # 永久令牌：二维码长期有效，扫码进入对应证书信息页
    token = make_qr_token(cert.id)
    path = url_for('public.detail', token=token)
    base = (current_app.config.get('SITE_BASE_URL') or '').rstrip('/')
    url = base + path if base else url_for('public.detail', token=token, _external=True)
    caption = current_app.config.get('QR_CAPTION') or (base + '/' if base else url)

    logo_path = os.path.join(current_app.static_folder, 'img', 'share.png')
    buf = generate_qr_png(url, logo_path=logo_path, caption=caption)
    image = 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode('ascii')
    return ok({'id': cert.id, 'name': cert.name, 'url': url,
               'caption': caption, 'image': image})


@admin_bp.route('/api/certificate/<int:cid>', methods=['POST'])
def api_certificate_update(cid):
    """更新单条证书（编辑弹窗用）"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    cert = Certificate.query.get_or_404(cid)
    payload = request.get_json(silent=True) or request.form
    data = {field: (payload.get(field) or '').strip() for field in FIELDS}

    errors = _validate(data)
    if not errors and _cert_no_exists(data['cert_no'], exclude_id=cid):
        errors.append('证书编号已存在：%s' % data['cert_no'])
    if errors:
        return fail('；'.join(errors), errors=errors)

    for field in FIELDS:
        setattr(cert, field, data[field])
    db.session.commit()
    return ok(_serialize(cert), message='保存成功')


@admin_bp.route('/certificates/batch_delete', methods=['POST'])
def certificates_batch_delete():
    """批量软删除"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)

    payload = request.get_json(silent=True) or {}
    ids = []
    for value in (payload.get('ids') or []):
        try:
            ids.append(int(value))
        except (TypeError, ValueError):
            continue
    if not ids:
        return fail('请选择要删除的证书')

    certs = Certificate.query.filter(Certificate.id.in_(ids),
                                     Certificate.is_deleted.is_(False)).all()
    now = datetime.now()
    for cert in certs:
        cert.is_deleted = True
        cert.deleted_at = now
    db.session.commit()
    return ok({'deleted': len(certs)}, message='已删除 %d 条' % len(certs))


@admin_bp.route('/certificates/import', methods=['POST'])
def certificates_import():
    """导入证书（Excel），逐行校验必填字段并做证书编号去重（返回 JSON）"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)

    file = request.files.get('file')
    if not file or not file.filename:
        return fail('请选择要导入的 Excel 文件')

    rows, fatal = parse_workbook(file.stream)
    if fatal:
        return fail(fatal)

    summary = {'total': len(rows), 'ok': 0, 'failed': 0, 'items': []}
    seen = set()
    for row in rows:
        cert_no = row.get('cert_no', '')
        name = row.get('name', '')
        errors = list(row.get('errors', []))

        if cert_no in seen:
            errors.append('文件内证书编号重复')
        elif _cert_no_exists(cert_no):
            errors.append('证书编号已存在')

        if errors:
            summary['failed'] += 1
            summary['items'].append({'row': row['_row'], 'name': name,
                                     'cert_no': cert_no, 'ok': False,
                                     'message': '；'.join(errors)})
            continue

        seen.add(cert_no)
        db.session.add(Certificate(**{f: row.get(f, '') for f in FIELDS}))
        summary['ok'] += 1
        summary['items'].append({'row': row['_row'], 'name': name,
                                 'cert_no': cert_no, 'ok': True,
                                 'message': '导入成功'})
    db.session.commit()
    return ok(summary, message='导入完成：成功 %d 条，失败 %d 条'
              % (summary['ok'], summary['failed']))


@admin_bp.route('/certificates/template')
@login_required
def certificates_template():
    buf = build_template()
    return send_file(buf, as_attachment=True,
                     download_name='证书导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@admin_bp.route('/certificates/export', methods=['GET', 'POST'])
@login_required
def certificates_export():
    """导出证书数据；传入 ids（GET 逗号分隔 / POST JSON）时只导出选中项"""
    raw_ids = []
    if request.method == 'POST':
        payload = request.get_json(silent=True) or {}
        raw_ids = payload.get('ids') or []
    elif request.args.get('ids'):
        raw_ids = request.args.get('ids').split(',')

    ids = []
    for value in raw_ids:
        try:
            ids.append(int(value))
        except (TypeError, ValueError):
            continue

    query = Certificate.query.filter(Certificate.is_deleted.is_(False))
    if ids:
        query = query.filter(Certificate.id.in_(ids))
    certs = query.order_by(Certificate.id.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '证书数据'
    headers = [c[0] for c in COLUMNS] + ['创建时间', '更新时间']
    header_fill = PatternFill('solid', fgColor='D9D9D9')
    header_font = Font(color='000000', bold=True)
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 18

    for r, cert in enumerate(certs, start=2):
        for c, field in enumerate(FIELDS, start=1):
            cell = ws.cell(row=r, column=c, value=getattr(cert, field, ''))
            cell.border = border
            if field in ('id_card', 'cert_no', 'birth_date', 'issue_date'):
                cell.number_format = '@'
        ws.cell(row=r, column=len(FIELDS) + 1,
                value=cert.created_at.strftime('%Y-%m-%d %H:%M:%S') if cert.created_at else '')
        ws.cell(row=r, column=len(FIELDS) + 2,
                value=cert.updated_at.strftime('%Y-%m-%d %H:%M:%S') if cert.updated_at else '')

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name='证书数据导出.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# --------------------------------------------------------------------------- #
# 考试报名管理
# --------------------------------------------------------------------------- #
@admin_bp.route('/registrations', methods=['GET'])
@login_required
def registrations():
    """报名管理页面（固定 HTML，数据由 /admin/api/registrations 填充）"""
    names = [row[0] for row in db.session.query(Registration.institution)
             .filter(Registration.institution.isnot(None), Registration.institution != '',
                     Registration.is_deleted.is_(False))
             .distinct().all()]
    return render_template('admin/registrations.html', institutions=sorted(names))


def _reg_serialize(reg):
    return reg.to_dict()


@admin_bp.route('/api/registrations')
def api_registrations():
    """报名列表接口（统一 JSON 格式：success + pagination + data）"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)

    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', current_app.config['PER_PAGE'], type=int)
    keyword = (request.args.get('q') or '').strip()
    subject = (request.args.get('subject') or '').strip()

    query = Registration.query.filter(Registration.is_deleted.is_(False))
    if keyword:
        like = '%%%s%%' % keyword
        query = query.filter(db.or_(Registration.name.like(like),
                                    Registration.id_card.like(like),
                                    Registration.subject.like(like)))
    if subject:
        query = query.filter(Registration.subject == subject)
    institution = (request.args.get('institution') or '').strip()
    if institution:
        query = query.filter(Registration.institution == institution)

    pagination = query.order_by(Registration.id.desc()).paginate(
        page=page, per_page=per_page, error_out=False)
    data = [_reg_serialize(r) for r in pagination.items]
    return ok(data, pagination=pagination_dict(pagination))


@admin_bp.route('/api/registration/<int:rid>', methods=['GET'])
def api_registration_get(rid):
    """获取单条报名（详情弹窗用）"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    reg = Registration.query.get_or_404(rid)
    return ok(_reg_serialize(reg))


@admin_bp.route('/api/registration/<int:rid>', methods=['POST'])
def api_registration_update(rid):
    """编辑单条报名"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    reg = Registration.query.get_or_404(rid)
    payload = request.form

    name = (payload.get('name') or '').strip()
    if not name:
        return fail('请填写考生姓名')

    id_card = (payload.get('id_card') or '').strip().upper()
    if id_card and not re.match(r'^(\d{15}|\d{17}[\dXx])$', id_card):
        return fail('身份证号格式不正确')

    reg.name = name
    reg.gender = (payload.get('gender') or '').strip()
    reg.ethnicity = (payload.get('ethnicity') or '').strip()
    reg.nationality = (payload.get('nationality') or '').strip()
    reg.id_card = id_card
    reg.subject = (payload.get('subject') or '').strip()
    reg.level = (payload.get('level') or '').strip()

    db.session.commit()
    return ok(_reg_serialize(reg), message='保存成功')


@admin_bp.route('/registration/<int:rid>/delete', methods=['POST'])
@login_required
def registration_delete(rid):
    reg = Registration.query.get_or_404(rid)
    reg.is_deleted = True
    reg.deleted_at = datetime.now()
    db.session.commit()
    flash('已删除', 'success')
    return redirect(request.referrer or url_for('admin.registrations'))


@admin_bp.route('/registrations/batch_delete', methods=['POST'])
def registrations_batch_delete():
    """批量软删除"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)

    payload = request.get_json(silent=True) or {}
    ids = []
    for value in (payload.get('ids') or []):
        try:
            ids.append(int(value))
        except (TypeError, ValueError):
            continue
    if not ids:
        return fail('请选择要删除的报名')

    regs = Registration.query.filter(Registration.id.in_(ids),
                                     Registration.is_deleted.is_(False)).all()
    now = datetime.now()
    for reg in regs:
        reg.is_deleted = True
        reg.deleted_at = now
    db.session.commit()
    return ok({'deleted': len(regs)}, message='已删除 %d 条' % len(regs))


@admin_bp.route('/registrations/export', methods=['GET', 'POST'])
@login_required
def registrations_export():
    """导出报名数据；传入 ids 时只导出选中项"""
    raw_ids = []
    if request.method == 'POST':
        payload = request.get_json(silent=True) or {}
        raw_ids = payload.get('ids') or []
    elif request.args.get('ids'):
        raw_ids = request.args.get('ids').split(',')

    ids = []
    for value in raw_ids:
        try:
            ids.append(int(value))
        except (TypeError, ValueError):
            continue

    query = Registration.query.filter(Registration.is_deleted.is_(False))
    if ids:
        query = query.filter(Registration.id.in_(ids))
    regs = query.order_by(Registration.id.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '报名数据'
    headers = ['所属机构', '考生姓名', '性别', '民族', '国籍', '身份证号', '报考科目', '报考级别',
               '报名时间']
    header_fill = PatternFill('solid', fgColor='D9D9D9')
    header_font = Font(color='000000', bold=True)
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 18

    for r, reg in enumerate(regs, start=2):
        values = [reg.institution, reg.name, reg.gender, reg.ethnicity, reg.nationality, reg.id_card,
                  reg.subject, reg.level,
                  reg.created_at.strftime('%Y-%m-%d %H:%M:%S') if reg.created_at else '']
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = border
            if c == 6:
                cell.number_format = '@'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name='报名数据导出.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# --------------------------------------------------------------------------- #
# 报名机构管理（按机构生成报名链接）
# --------------------------------------------------------------------------- #
def _signup_link(code):
    """机构报名链接：按当前访问域名生成（本地即 127.0.0.1，线上即线上域名）。"""
    return url_for('public.signup', i=code, _external=True)


def _inst_serialize(inst):
    data = inst.to_dict()
    data['link'] = _signup_link(inst.code)
    data['count'] = Registration.query.filter(
        Registration.institution_id == inst.id, Registration.is_deleted.is_(False)).count()
    return data


@admin_bp.route('/institutions', methods=['GET'])
@login_required
def institutions():
    """机构管理页面（固定 HTML，数据由 /admin/api/institutions 填充）"""
    return render_template('admin/institutions.html')


@admin_bp.route('/api/institutions')
def api_institutions():
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', current_app.config['PER_PAGE'], type=int)
    keyword = (request.args.get('q') or '').strip()

    query = Institution.query.filter(Institution.is_deleted.is_(False))
    if keyword:
        query = query.filter(Institution.name.like('%%%s%%' % keyword))
    pagination = query.order_by(Institution.id.desc()).paginate(
        page=page, per_page=per_page, error_out=False)
    data = [_inst_serialize(i) for i in pagination.items]
    return ok(data, pagination=pagination_dict(pagination))


@admin_bp.route('/api/institution', methods=['POST'])
def api_institution_create():
    """新增机构并生成报名链接"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    payload = request.get_json(silent=True) or request.form
    name = (payload.get('name') or '').strip()
    if not name:
        return fail('请填写机构名称')
    if len(name) > 100:
        return fail('机构名称过长')
    if Institution.query.filter(Institution.name == name, Institution.is_deleted.is_(False)).first():
        return fail('该机构已存在')

    import secrets
    code = secrets.token_urlsafe(6).replace('-', 'A').replace('_', 'B')[:10]
    while Institution.query.filter_by(code=code).first():
        code = secrets.token_urlsafe(6).replace('-', 'A').replace('_', 'B')[:10]
    inst = Institution(name=name, code=code)
    db.session.add(inst)
    db.session.commit()
    return ok(_inst_serialize(inst), message='机构已创建')


@admin_bp.route('/institution/<int:iid>/delete', methods=['POST'])
@login_required
def institution_delete(iid):
    inst = Institution.query.get_or_404(iid)
    inst.is_deleted = True
    db.session.commit()
    flash('已删除', 'success')
    return redirect(request.referrer or url_for('admin.institutions'))


@admin_bp.route('/api/institution/<int:iid>/qrcode')
def api_institution_qrcode(iid):
    """生成机构报名链接的二维码（中间带 logo、底部带链接），返回链接与 data URI 图片"""
    if not current_user.is_authenticated:
        return fail('未登录或登录已过期', status=401)
    inst = Institution.query.get_or_404(iid)
    import base64
    url = _signup_link(inst.code)
    caption = (current_app.config.get('SITE_BASE_URL') or '').rstrip('/') or url
    logo_path = os.path.join(current_app.static_folder, 'img', 'share.png')
    buf = generate_qr_png(url, logo_path=logo_path, caption=caption)
    image = 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode('ascii')
    return ok({'id': inst.id, 'name': inst.name, 'url': url, 'caption': caption, 'image': image})
